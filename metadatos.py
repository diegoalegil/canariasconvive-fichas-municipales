"""Fuentes institucionales leídas del índice del libro, sin créditos personales."""
from openpyxl import load_workbook


def fuentes_indicadores(ruta, fichas):
    libro = load_workbook(ruta, data_only=False)
    hoja = libro['INDEX-F']

    def enlaces(filas):
        salida = []
        for fila in filas:
            celda = hoja.cell(fila, 5)
            if not celda.hyperlink:
                raise ValueError(f'Falta el enlace estadístico en INDEX-F!E{fila}')
            revisado = hoja.cell(fila, 11).value
            salida.append({'organismo': hoja.cell(fila, 10).value,
                           'url': celda.hyperlink.target,
                           'desde': hoja.cell(fila, 7).value,
                           'hasta': hoja.cell(fila, 8).value,
                           'revision': revisado.date().isoformat() if revisado else None})
        return salida

    poblacion = enlaces([2, 3])
    nacimiento = enlaces([5, 6])
    reciente = enlaces([3])
    anio = max(f['anio'] for f in fichas)
    serie = f"{min(f['evolucion']['anios'][0] for f in fichas)}–{anio}"
    base = f"{min(f['evolucion']['anio_base'] for f in fichas)}–{anio}"
    comp = fichas[0]['componentes']
    def periodo(clave):
        fechas = [a for f in fichas for a, v in zip(f['componentes']['anios'], f['componentes'][clave]) if v is not None]
        return f'{min(fechas)}–{max(fechas)}'
    defs = {
        'poblacion': ('Habitantes', f'1 de enero de {anio}', poblacion, 'Recuento de residentes.'),
        'sexo': ('Mujeres y hombres', f'1 de enero de {anio}', reciente, 'Recuento y porcentaje sobre la población total del municipio.'),
        'edad': ('Edad media', f'1 de enero de {anio}', reciente, 'Aproximación calculada con los grupos quinquenales: marcas de clase 2,5; 7,5; …; 97,5 años y 102 años para el grupo de 100 o más.'),
        'evolucion': ('Evolución y variación acumulada', serie, poblacion, 'La variación acumulada compara la población final con la inicial del periodo indicado en la ficha.'),
        'tvma': ('Variación media anual', base, poblacion, 'Tasa anual equivalente: [(población final / población inicial)^(1/n) − 1] × 100. n es el número de años transcurridos.'),
        'extranjero': ('Origen extranjero', f'2000–{anio}', nacimiento, 'Personas nacidas fuera de España, con independencia de su nacionalidad. Porcentaje sobre el total de residentes.'),
        'piramide': ('Estructura de la población', f'1 de enero de {anio}', reciente + enlaces([6]), 'Municipio y Canarias: cada uno sobre su población total. Por lugar de nacimiento: cada grupo sobre su propio total, sumando hombres y mujeres.'),
        'nacimiento': ('Lugar de nacimiento', f'1 de enero de {anio}', enlaces([6]), 'Porcentaje nacido en Canarias, en el resto de España y en el extranjero.'),
        'vegetativo': ('Crecimiento vegetativo', periodo('vegetativo'), enlaces([18]), 'Nacimientos menos defunciones en el mismo año.'),
        'migratorio': ('Saldo migratorio', periodo('migratorio'), enlaces([19, 20]), 'Entradas menos salidas por cambio de residencia en el mismo año.'),
        'rankings': ('El municipio en su entorno', f'1 de enero de {anio}', poblacion, 'Posición por habitantes y porcentaje que representan en cada ámbito territorial.'),
    }
    for clave, cod, texto in [
        ('envejecimiento', 'C10', 'Personas de 65 años o más por cada persona menor de 15.'),
        ('juventud', 'C11', 'Menores de 15 por cada cien personas de 15 a 64. La base es la población de 15 a 64 años.'),
        ('dependencia', 'C17', 'Menores de 15 y mayores de 64 por cada cien personas de 15 a 64.'),
        ('reemplazo', 'C14', 'Personas de 15 a 19 por cada cien de 60 a 64.')]:
        defs[clave] = (fichas[0]['indices'][cod]['etiqueta'], str(fichas[0]['indices'][cod]['anio']), reciente, texto)
    libro.close()
    return {k: {'titulo': t, 'periodo': p, 'enlaces': e, 'nota': n} for k, (t, p, e, n) in defs.items()}
