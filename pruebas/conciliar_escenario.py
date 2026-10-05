#!/usr/bin/env python3
"""Conciliación del escenario 2035 (web/datos/escenario/) contra el libro de la
proyección de Pedro, PROYECCION_HP_2035.xlsx, por un camino distinto del de
exportar_escenario.py: los índices, la población y la variación media anual,
contra los que calcula el propio libro en su hoja 21_INDICES_TVMA (municipios,
islas, provincias y Canarias), y cada barra de las pirámides de 2025 y 2035,
contra las hojas POB_2025 y 25_VALORES_2035 sumadas aquí por territorio. Y,
para cada cifra, que el decimal que escribe la página es el del valor del
libro redondeado una sola vez (README, «Redondeo una sola vez»).

Necesita el libro en ~/Downloads (o en la ruta que se pase como argumento) y
openpyxl; si no está el libro, se omite avisando, como conciliar_excel.py."""
import json
import sys
from collections import Counter
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
RUTA = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / "Downloads" / "PROYECCION_HP_2035.xlsx"
if not RUTA.exists():
    print(f"omitida · no está el libro de la proyección en {RUTA}")
    sys.exit(0)
try:
    import openpyxl  # noqa: E402
except ImportError:
    print("FALLA · está el libro pero falta openpyxl: ejecutar con el python3 que lo tenga (/usr/bin/python3)")
    sys.exit(1)

W = openpyxl.load_workbook(RUTA, read_only=True, data_only=True)
DATOS = RAIZ / "web/datos"
INDICE = json.loads((DATOS / "indice.json").read_text(encoding="utf-8"))
ERR, CUENTA = [], Counter()


def check(ok, donde, detalle=""):
    CUENTA[donde.split(":")[0]] += 1
    if not ok:
        ERR.append(f"{donde} {detalle}")


def mostrado(v, dec):
    """Lo que escribe la web con `nf(v, dec)`: half-expand sobre el valor decimal."""
    return str(Decimal(repr(float(v))).quantize(Decimal(1).scaleb(-dec), rounding=ROUND_HALF_UP))


# Los ficheros del escenario, por (ámbito del libro, nombre).
ESC = {}
for m in INDICE["municipios"]:
    ESC[("Municipio", m["nombre"])] = json.loads((DATOS / f"escenario/mun/{m['codmun']}.json").read_text(encoding="utf-8"))
for i in INDICE["islas_resumen"]:
    ESC[("Isla", i["nombre"])] = json.loads((DATOS / f"escenario/isla/{i['slug']}.json").read_text(encoding="utf-8"))
for p in INDICE["provincias"]:
    ESC[("Provincia", p["nombre"])] = json.loads((DATOS / f"escenario/provincia/{p['slug']}.json").read_text(encoding="utf-8"))
ESC[("Canarias", "Canarias")] = json.loads((DATOS / "escenario/canarias.json").read_text(encoding="utf-8"))

# ---- 21_INDICES_TVMA: lo que calcula el libro, fila a fila.
filas = list(W["21_INDICES_TVMA"].iter_rows(values_only=True))
cab = list(filas[0])
col = {c: k for k, c in enumerate(cab)}
NOMBRES = {"C10": "Envejecimiento", "C11": "Juventud", "C17": "Dependencia", "C14": "Reemplazo laboral"}
vistos = set()
for f in filas[1:]:
    if not f or f[0] is None:
        continue
    clave = (str(f[0]).strip(), str(f[1]).strip())
    e = ESC.get(clave)
    check(e is not None, "territorio:", f"{clave} está en el libro y no en la web")
    if e is None:
        continue
    vistos.add(clave)
    b, h = e["base"], e["horizonte"]
    for k, anio in enumerate(e["proyectado"]["anios"]):
        check(e["proyectado"]["valores"][k] == round(f[col[f"Población {anio}"]]), f"población:{clave[1]} {anio}",
              f"{e['proyectado']['valores'][k]} frente a {f[col[f'Población {anio}']]}")
    for cod, nombre in NOMBRES.items():
        mult = 100 if cod == "C11" else 1   # el libro guarda la juventud en tanto por uno
        for anio, valor in ((b, e["indices"][cod]["base"]), (h, e["indices"][cod]["horizonte"])):
            libro = f[col[f"{nombre} {anio}"]] * mult
            check(abs(valor - libro) < 1e-5, f"índice:{clave[1]} {cod} {anio}", f"{valor} frente a {libro}")
            dec = e["indices"][cod]["decimales"]
            check(mostrado(valor, dec) == mostrado(libro, dec), f"decimal:{clave[1]} {cod} {anio}", f"{mostrado(valor, dec)} frente a {mostrado(libro, dec)}")
    tvma = f[col[f"TVMA {b}-{h} (%)"]]
    check(abs(e["tvma"]["filas"][0]["proyectada"] - tvma) < 1e-5, f"tvma:{clave[1]}", f"{e['tvma']['filas'][0]['proyectada']} frente a {tvma}")
    check(mostrado(e["tvma"]["filas"][0]["proyectada"], 1) == mostrado(tvma, 1), f"decimal:{clave[1]} tvma", f"frente a {tvma}")
check(vistos == set(ESC), "territorio:", f"faltan en el libro {sorted(set(ESC) - vistos)}")

# ---- POB_2025 y 25_VALORES_2035: cada barra de las dos pirámides, en porcentaje sobre el total del territorio.
def hoja(nombre):
    F = list(W[nombre].iter_rows(values_only=True))
    columnas, actual = {}, None
    for c in range(2, len(F[1])):
        if F[0][c] is not None:
            actual = str(F[0][c]).strip()
        if F[1][c]:
            columnas[(actual, "hombres" if str(F[1][c]).lower().startswith("hombre") else "mujeres")] = c
    return F, columnas


HOJAS = {"base": hoja(f"POB_{ESC[('Canarias', 'Canarias')]['base']}"), "horizonte": hoja(f"25_VALORES_{ESC[('Canarias', 'Canarias')]['horizonte']}")}
islas_de = {p["nombre"]: [i["nombre"] for i in INDICE["islas_resumen"] if i["slug"] in p["islas"]] for p in INDICE["provincias"]}
muns_de = dict(INDICE["islas"])
for (ambito, nombre), e in ESC.items():
    muns = ([nombre] if ambito == "Municipio" else muns_de[nombre] if ambito == "Isla"
            else [n for i in islas_de[nombre] for n in muns_de[i]] if ambito == "Provincia" else [m["nombre"] for m in INDICE["municipios"]])
    for anio, (F, columnas) in HOJAS.items():
        lados = {lado: [sum(F[2 + g][columnas[(n, lado)]] for n in muns) for g in range(21)] for lado in ("hombres", "mujeres")}
        total = sum(lados["hombres"]) + sum(lados["mujeres"])
        for lado, v in lados.items():
            pct = [x / total * 100 for x in v[:17] + [sum(v[17:])]]
            for g, (a, c) in enumerate(zip(e["piramide"][anio][lado], pct)):
                donde = f"{nombre} {anio} {lado} {e['piramide']['edades'][g]}"
                check(abs(a - c) < 1e-5, f"pirámide:{donde}", f"{a} frente a {c}")
                check(mostrado(a, 1) == mostrado(c, 1), f"decimal:{donde}", f"{mostrado(a, 1)} frente a {mostrado(c, 1)} ({c})")

if ERR:
    print(f"FALLA · {len(ERR)} diferencia(s) con {RUTA.name}:")
    for x in ERR[:40]:
        print(" -", x)
    sys.exit(1)
print(f"ok · escenario 2035 contra {RUTA.name}: " + ", ".join(f"{n} {k}" for k, n in sorted(CUENTA.items()))
      + f" ({sum(CUENTA.values())} comparaciones)")
